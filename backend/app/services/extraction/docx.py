"""DOCX text extraction via python-docx (REQ-1.3).

DOCX carries explicit paragraph styles, so heading detection here is a style
lookup rather than the font-size heuristic the PDF path needs.

Tables are included: Indian vendor agreements routinely put payment schedules
and service levels in tables, and dropping them would lose exactly the clauses
the rubric most wants to inspect.
"""

from __future__ import annotations

import io
import logging

from app.services.extraction.base import (
    CorruptFileError,
    ExtractedBlock,
    ExtractedDocument,
    NoExtractableTextError,
)

logger = logging.getLogger(__name__)

MIN_MEANINGFUL_CHARS = 200


def extract_docx(data: bytes) -> ExtractedDocument:
    import docx  # imported lazily so the module imports without the dep
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise CorruptFileError(str(exc)) from exc

    parts: list[str] = []
    blocks: list[ExtractedBlock] = []
    cursor = 0

    for item in _iter_body(document, Paragraph, Table):
        if isinstance(item, Paragraph):
            text = item.text.strip()
            if not text:
                continue
            is_heading = _is_heading_style(item)
        else:
            text = _render_table(item)
            if not text:
                continue
            is_heading = False

        start = cursor
        parts.append(text)
        cursor += len(text) + 1

        blocks.append(
            ExtractedBlock(
                text=text,
                char_start=start,
                char_end=start + len(text),
                is_heading=is_heading,
            )
        )

    assembled = "\n".join(parts)

    if len(assembled.strip()) < MIN_MEANINGFUL_CHARS:
        raise NoExtractableTextError(f"only {len(assembled.strip())} characters extracted")

    return ExtractedDocument(text=assembled, blocks=blocks)


def _iter_body(document: object, paragraph_cls: type, table_cls: type):
    """Yield paragraphs and tables in document order.

    python-docx exposes ``.paragraphs`` and ``.tables`` as separate lists,
    which loses their relative order. Walking the underlying XML body keeps a
    payment-terms table attached to the heading that introduces it.
    """
    body = document.element.body  # type: ignore[attr-defined]
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            yield paragraph_cls(child, document)  # type: ignore[arg-type]
        elif tag == "tbl":
            yield table_cls(child, document)  # type: ignore[arg-type]


def _is_heading_style(paragraph: object) -> bool:
    style_name = (getattr(paragraph.style, "name", "") or "").lower()  # type: ignore[attr-defined]
    if "heading" in style_name or style_name in {"title", "subtitle"}:
        return True

    text = paragraph.text.strip()  # type: ignore[attr-defined]
    if not text or len(text) > 80:
        return False

    # Fall back to the same conventions as the PDF path for documents that
    # were written with direct formatting instead of styles.
    runs = getattr(paragraph, "runs", ())
    if runs and all(r.bold for r in runs if r.text.strip()):
        return True
    letters = [c for c in text if c.isalpha()]
    return bool(letters) and all(c.isupper() for c in letters)


def _render_table(table: object) -> str:
    rows: list[str] = []
    for row in table.rows:  # type: ignore[attr-defined]
        cells = [c.text.strip().replace("\n", " ") for c in row.cells]
        # De-duplicate horizontally merged cells, which python-docx repeats.
        deduped: list[str] = []
        for cell in cells:
            if not deduped or deduped[-1] != cell:
                deduped.append(cell)
        line = " | ".join(c for c in deduped if c)
        if line:
            rows.append(line)
    return "\n".join(rows)
