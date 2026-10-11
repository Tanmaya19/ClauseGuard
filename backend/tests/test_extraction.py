"""Extraction tests (REQ-1.1, REQ-1.3, REQ-1.4)."""

from __future__ import annotations

import pytest

from app.services.extraction import (
    CorruptFileError,
    NoExtractableTextError,
    UnsupportedFormatError,
    detect_format,
    extract,
)
from tests.fixtures import (
    RISKY_AGREEMENT,
    make_docx,
    make_image_only_pdf,
    make_legacy_doc,
    make_pdf,
)


def test_detects_pdf_from_magic_bytes() -> None:
    assert detect_format(make_pdf(), "anything.txt") == "pdf"


def test_detects_docx_from_container_contents() -> None:
    # Extension deliberately wrong: content must win, because the extension is
    # user-supplied.
    assert detect_format(make_docx(), "renamed.pdf") == "docx"


def test_rejects_legacy_doc_with_a_specific_message() -> None:
    with pytest.raises(UnsupportedFormatError, match="save it as .docx"):
        detect_format(make_legacy_doc(), "old.doc")


def test_rejects_a_zip_that_is_not_a_word_document() -> None:
    import io
    import zipfile

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("xl/workbook.xml", "<workbook/>")

    with pytest.raises(UnsupportedFormatError):
        detect_format(buffer.getvalue(), "sheet.xlsx")


def test_rejects_empty_upload() -> None:
    with pytest.raises(CorruptFileError):
        extract(b"", "empty.pdf")


def test_rejects_arbitrary_bytes() -> None:
    with pytest.raises(UnsupportedFormatError):
        extract(b"just some text, not a document at all" * 20, "notes.txt")


@pytest.mark.parametrize("builder", [make_pdf, make_docx])
def test_extracts_full_text(builder) -> None:
    document = extract(builder(), "contract")

    assert not document.is_empty
    # Spot-check content from the first and last clause, so a truncating bug
    # cannot pass.
    assert "Sunrise Components" in document.text
    assert "registered post" in document.text
    assert len(document.blocks) > len(RISKY_AGREEMENT)


@pytest.mark.parametrize("builder", [make_pdf, make_docx])
def test_block_offsets_point_at_the_assembled_text(builder) -> None:
    """Offsets are what the UI uses to highlight the source span, so an
    off-by-one here would mis-highlight every clause."""
    document = extract(builder(), "contract")

    for block in document.blocks:
        assert document.text[block.char_start : block.char_end] == block.text


@pytest.mark.parametrize("builder", [make_pdf, make_docx])
def test_identifies_headings(builder) -> None:
    document = extract(builder(), "contract")
    headings = [b.text for b in document.blocks if b.is_heading]

    assert any("Indemnity" in h for h in headings)
    assert any("Governing Law" in h for h in headings)


def test_image_only_pdf_is_reported_as_unreadable() -> None:
    """REQ-1.4: OCR is out of scope (TBD-7), so this is terminal."""
    with pytest.raises(NoExtractableTextError) as exc:
        extract(make_image_only_pdf(), "scan.pdf")

    assert exc.value.error_code == "no_extractable_text"
    # The message reaches the user, so it must not be technical.
    assert "scan" in exc.value.user_message.lower()


def test_pdf_and_docx_of_the_same_contract_agree() -> None:
    """Guards the pipeline against format-dependent behaviour.

    Not byte-for-byte equality: a PDF is laid out in *lines*, so its blocks
    break mid-sentence where the renderer wrapped them, while DOCX blocks are
    whole paragraphs. The invariant that matters is that the words are the
    same, and that segmentation therefore sees the same contract either way —
    which ``tests/test_segmentation.py`` asserts on the clause level.
    """
    pdf_text = extract(make_pdf(), "a.pdf").text
    docx_text = extract(make_docx(), "a.docx").text

    assert " ".join(pdf_text.split()) == " ".join(docx_text.split())


def test_docx_tables_are_included() -> None:
    """Indian vendor agreements routinely put payment schedules in tables;
    dropping them would lose the clauses the rubric most wants to see."""
    import docx

    document = docx.Document()
    document.add_paragraph("1. Payment Terms")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Milestone"
    table.cell(0, 1).text = "Amount"
    table.cell(1, 0).text = "On delivery"
    table.cell(1, 1).text = "Rs. 4,50,000 within 90 days"
    document.add_paragraph("x" * 300)  # clear the minimum-content threshold

    import io

    buffer = io.BytesIO()
    document.save(buffer)

    assert "4,50,000" in extract(buffer.getvalue(), "t.docx").text
