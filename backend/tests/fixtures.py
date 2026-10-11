"""Synthetic contract fixtures, built in-memory as real PDF and DOCX bytes.

Generated rather than committed as binaries so the test suite stays readable:
the clause text that a test asserts on is visible in this file, not hidden
inside a checked-in blob.

The sample agreement is deliberately seeded with specific risks so the rubric
has something to find — see ``EXPECTED_RISKS`` for the mapping from clause to
the rule that should fire.
"""

from __future__ import annotations

import io

# A short Indian vendor agreement. Numbered clauses, so the numbering strategy
# should win. Each entry is (heading-and-number, body).
RISKY_AGREEMENT: list[tuple[str, str]] = [
    (
        "1. Parties and Definitions",
        "This Agreement is made on 1 April 2026 between Sunrise Components Private "
        "Limited, a company incorporated under the Companies Act, 2013, having its "
        "registered office at Coimbatore, Tamil Nadu (hereinafter the \"Supplier\"), "
        "and Meridian Industries Limited, having its registered office at Mumbai, "
        "Maharashtra (hereinafter the \"Customer\"). \"Goods\" shall mean the "
        "components described in Schedule 1.",
    ),
    (
        "2. Term",
        "This Agreement shall commence on the Effective Date and shall remain in "
        "force for an initial period of twenty-four months unless terminated "
        "earlier in accordance with its terms.",
    ),
    (
        "3. Renewal",
        "This Agreement shall automatically renew for successive periods of twelve "
        "months at the end of the initial term. The Agreement shall be deemed "
        "renewed unless written notice is given.",
    ),
    (
        "4. Payment Terms",
        "The Customer shall pay each invoice raised by the Supplier within ninety "
        "(90) days from the date of receipt of the invoice. All prices are "
        "exclusive of GST, which shall be charged at the applicable rate. The "
        "Customer reserves the right to revise the price of any Goods upon "
        "intimation to the Supplier.",
    ),
    (
        "5. Indemnity",
        "The Supplier shall indemnify, defend and hold harmless the Customer "
        "against any and all claims, losses, damages, liabilities, costs and "
        "expenses, including loss of profits and consequential damages, arising "
        "out of or in connection with the supply of the Goods.",
    ),
    (
        "6. Limitation of Liability",
        "The Supplier shall be liable for all direct and indirect losses suffered "
        "by the Customer arising from any defect in the Goods or any delay in "
        "delivery.",
    ),
    (
        "7. Termination for Convenience",
        "The Customer may terminate this Agreement at any time without assigning "
        "any reason and without liability to the Supplier.",
    ),
    (
        "8. Termination for Cause",
        "The Customer may terminate this Agreement immediately upon a material "
        "breach by the Supplier of any provision of this Agreement.",
    ),
    (
        "9. Confidentiality",
        "The Supplier shall not disclose any Confidential Information of the "
        "Customer to any third party. This obligation shall continue in perpetuity "
        "notwithstanding the expiry or termination of this Agreement.",
    ),
    (
        "10. Intellectual Property",
        "The Supplier hereby assigns to the Customer all right, title and interest "
        "in any and all intellectual property created, developed or used in "
        "connection with the performance of this Agreement.",
    ),
    (
        "11. Exclusivity",
        "The Supplier shall supply the Goods exclusively to the Customer and shall "
        "not supply to any other person any goods of a similar description during "
        "the term of this Agreement.",
    ),
    (
        "12. Assignment",
        "The Supplier shall not assign, novate or otherwise transfer its rights or "
        "obligations under this Agreement without the prior written consent of the "
        "Customer.",
    ),
    (
        "13. Audit",
        "The Customer may at any time and without notice inspect the books and "
        "records of the Supplier relating to this Agreement and may require "
        "unrestricted access to the Supplier's premises.",
    ),
    (
        "14. Governing Law and Jurisdiction",
        "This Agreement shall be governed by and construed in accordance with the "
        "laws of Singapore, and the courts at Singapore shall have exclusive "
        "jurisdiction over any dispute arising hereunder.",
    ),
    (
        "15. Notices",
        "Any notice under this Agreement shall be in writing and sent by registered "
        "post to the address of the receiving party set out above.",
    ),
]

# Rules we expect to fire, by clause heading. Used by the rubric tests, and the
# reason this fixture is shaped the way it is.
EXPECTED_RISKS: dict[str, set[str]] = {
    "3. Renewal": {"auto_renew_without_notice_window"},
    "4. Payment Terms": {"payment_term_over_45_days", "unilateral_price_variation"},
    "5. Indemnity": {"indemnity_covers_indirect_losses", "uncapped_liability"},
    "6. Limitation of Liability": {"uncapped_liability"},
    "7. Termination for Convenience": {"termination_without_notice_period"},
    "8. Termination for Cause": {"no_cure_period"},
    "9. Confidentiality": {"perpetual_confidentiality"},
    "10. Intellectual Property": {"broad_ip_assignment"},
    "11. Exclusivity": {"exclusivity_without_commitment"},
    "13. Audit": {"unlimited_audit_rights"},
    "14. Governing Law and Jurisdiction": {"foreign_governing_law"},
}

# Clause types absent from the fixture, so the contract-level rules fire.
EXPECTED_CONTRACT_RISKS = {"no_force_majeure_clause"}


def _plain_text(clauses: list[tuple[str, str]] | None = None) -> str:
    clauses = clauses if clauses is not None else RISKY_AGREEMENT
    return "\n".join(f"{heading}\n{body}" for heading, body in clauses)


def make_docx(clauses: list[tuple[str, str]] | None = None) -> bytes:
    """A real .docx with Heading 2 styles on clause headings."""
    import docx

    clauses = clauses if clauses is not None else RISKY_AGREEMENT
    document = docx.Document()
    document.add_heading("SUPPLY AGREEMENT", level=1)

    for heading, body in clauses:
        document.add_heading(heading, level=2)
        document.add_paragraph(body)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pdf(clauses: list[tuple[str, str]] | None = None) -> bytes:
    """A real, text-extractable PDF with larger/bold headings."""
    from reportlab.lib.enums import TA_JUSTIFY
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    clauses = clauses if clauses is not None else RISKY_AGREEMENT
    styles = getSampleStyleSheet()
    heading = ParagraphStyle(
        "ClauseHeading", parent=styles["Heading2"], fontSize=13, spaceAfter=4
    )
    body = ParagraphStyle(
        "ClauseBody", parent=styles["BodyText"], fontSize=10, alignment=TA_JUSTIFY
    )

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=20 * mm, bottomMargin=20 * mm)

    flow = [Paragraph("SUPPLY AGREEMENT", styles["Title"]), Spacer(1, 6 * mm)]
    for clause_heading, clause_body in clauses:
        flow.append(Paragraph(clause_heading, heading))
        flow.append(Paragraph(clause_body, body))
        flow.append(Spacer(1, 3 * mm))

    doc.build(flow)
    return buffer.getvalue()


def make_image_only_pdf() -> bytes:
    """A PDF containing a drawing and no extractable text (REQ-1.4)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=A4)
    # Vector shapes only — no text operators at all.
    for i in range(40):
        c.rect(40 + i * 3, 40 + i * 5, 120, 30, stroke=1, fill=0)
    c.showPage()
    c.save()
    return buffer.getvalue()


def make_unnumbered_docx() -> bytes:
    """Paragraphs with no numbering and no heading styles (REQ-2.3 fallback)."""
    import docx

    document = docx.Document()
    for _, body in RISKY_AGREEMENT:
        document.add_paragraph(body)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_legacy_doc() -> bytes:
    """OLE2 header, as a real .doc would start (REQ-1.1 rejection path)."""
    return b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512
